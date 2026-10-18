import io
import json
import logging
import os
import sys
import threading
from datetime import datetime
from http.server import BaseHTTPRequestHandler, HTTPServer

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# Garante suporte a emojis no terminal Windows
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

from dotenv import load_dotenv
from telegram import Update
from telegram.ext import (
    Application,
    CommandHandler,
    ContextTypes,
    MessageHandler,
    filters,
)
from google import genai
from google.genai import types

# Importa o módulo do banco de dados (Supabase + fallback local)
from database import carregar_todos_dados, deletar_usuario, salvar_todos_dados, salvar_usuario

# Diretório base do projeto
BASE_DIR = os.path.dirname(os.path.abspath(__file__))

# Carrega variáveis do .env
load_dotenv(os.path.join(BASE_DIR, ".env"))

logging.basicConfig(format="%(asctime)s - %(name)s - %(levelname)s - %(message)s", level=logging.INFO)

# 🔑 CONFIGURAÇÕES DE TOKENS
TELEGRAM_TOKEN = os.getenv("TELEGRAM_TOKEN")
if TELEGRAM_TOKEN:
    TELEGRAM_TOKEN = TELEGRAM_TOKEN.strip().replace("\n", "").replace("\r", "").replace(" ", "")

GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")
if GEMINI_API_KEY:
    GEMINI_API_KEY = GEMINI_API_KEY.strip().replace("\n", "").replace("\r", "").replace(" ", "")

ai_client = genai.Client(api_key=GEMINI_API_KEY) if GEMINI_API_KEY else None

# 🌐 SERVIDOR WEB PARA MONITORAMENTO E KEEPALIVE NA NUVEM (RENDER)
class KeepAliveHandler(BaseHTTPRequestHandler):
    def do_GET(self):
        self.send_response(200)
        self.send_header("Content-type", "application/json; charset=utf-8")
        self.end_headers()
        self.wfile.write(b'{"status": "online", "service": "Porquinho IA 24/7"}')

    def log_message(self, format, *args):
        return

def iniciar_servidor_web():
    porta = int(os.getenv("PORT", 10000))
    try:
        servidor = HTTPServer(("0.0.0.0", porta), KeepAliveHandler)
        thread = threading.Thread(target=servidor.serve_forever, daemon=True)
        thread.start()
        logging.info(f"🌐 Servidor de monitoramento (Keep-Alive) rodando na porta {porta}")
    except Exception as e:
        logging.warning(f"⚠️ Servidor web na porta {porta} nao pode ser iniciado: {e}")

# 🛠️ HELPER CENTRALIZADO: REGISTRAR E PROCESSAR GASTO
def registrar_novo_gasto(todos_dados, user_id, item, valor_gasto, categoria, canal="📝"):
    dados_usuario = todos_dados[user_id]
    data_atual = datetime.now().strftime("%d/%m/%Y")

    dados_usuario["saldo_atual"] -= valor_gasto
    dados_usuario["gastos"].append({
        "item": item,
        "valor": valor_gasto,
        "categoria": categoria,
        "data": data_atual
    })

    todos_dados[user_id] = dados_usuario
    salvar_usuario(user_id, dados_usuario)

    saldo_restante = dados_usuario["saldo_atual"]
    orcamento = dados_usuario["orcamento_inicial"]

    alerta_meta = ""
    if saldo_restante <= (orcamento * 0.5) and saldo_restante > (orcamento * 0.2):
        alerta_meta = "\n\n⚠️ **Atenção:** Você já gastou mais de 50% do seu limite disponível!"
    elif saldo_restante <= (orcamento * 0.2) and saldo_restante > 0:
        alerta_meta = "\n\n🚨 **ALERTA CRÍTICO:** Você tem menos de 20% do seu saldo restante!"
    elif saldo_restante <= 0:
        alerta_meta = "\n\n🛑 **SALDO ESTOURADO:** Seu orçamento zerou ou ficou negativo!"

    resposta_formatada = (
        f"✅ **Gasto Registrado via {canal}!**\n"
        f"📝 {item} ({categoria})\n"
        f"💰 **R$ {valor_gasto:.2f}**\n"
        f"⚙️ {data_atual}\n\n"
        f"📉 **Saldo Restante:** R$ {saldo_restante:.2f}"
        f"{alerta_meta}"
    )
    return resposta_formatada

# 🚀 COMANDO /start
async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = str(update.effective_user.id)
    user_name = update.effective_user.first_name
    todos_dados = carregar_todos_dados()

    if user_id in todos_dados:
        saldo = todos_dados[user_id]["saldo_atual"]
        await update.message.reply_text(
            f"Olá de volta, {user_name}! 🐷\n"
            f"Seu saldo atual é de **R$ {saldo:.2f}**.\n\n"
            "📌 **Comandos úteis:**\n"
            "📊 /extrato — Extrato com gráfico de gastos\n"
            "↩️ /desfazer — Apaga o último gasto registrado\n"
            "📑 /planilha — Baixa o extrato em Excel (.xlsx)\n"
            "💡 /ajuda <pergunta> — Consultoria financeira com IA\n"
            "🔄 /reset — Começar um novo mês\n\n"
            "✨ **Super Poderes do Porquinho:**\n"
            "• Digite qualquer gasto em texto (ex: *45 padaria*)\n"
            "• Envie **áudios** falando o que comprou!\n"
            "• Envie **fotos** de cupons, notas fiscais ou comprovantes de Pix!"
        )
    else:
        context.user_data["configurando_renda"] = True
        await update.message.reply_text(
            f"Olá, {user_name}! Bem-vindo ao Porquinho IA. 🐷\n\n"
            "Digite qual será o seu **orçamento ou limite de gastos** para este mês:\n"
            "(Exemplo: 704 ou 1500)"
        )

# 📊 COMANDO /extrato (COM GRÁFICO DE PIZZA)
async def extrato(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = str(update.effective_user.id)
    todos_dados = carregar_todos_dados()

    if user_id not in todos_dados or not todos_dados[user_id]["gastos"]:
        await update.message.reply_text("📭 Você ainda não registrou nenhum gasto este mês!")
        return

    dados_user = todos_dados[user_id]

    # Agrupa gastos por categoria
    categorias = {}
    total_gasto = 0.0

    for g in dados_user["gastos"]:
        cat = g["categoria"]
        if cat not in categorias:
            categorias[cat] = 0.0
        categorias[cat] += g["valor"]
        total_gasto += g["valor"]

    # Monta a legenda do extrato
    mensagem = "📊 **Seu Extrato do Mês:**\n\n"
    for cat, total_cat in sorted(categorias.items(), key=lambda x: x[1], reverse=True):
        porcentagem = (total_cat / total_gasto * 100) if total_gasto > 0 else 0
        mensagem += f"📂 *{cat}:* R$ {total_cat:.2f} ({porcentagem:.1f}%)\n"

    mensagem += "───────────────────\n"
    mensagem += f"💸 **Total Gasto:** R$ {total_gasto:.2f}\n"
    mensagem += f"💰 **Saldo Restante:** R$ {dados_user['saldo_atual']:.2f}\n"
    mensagem += "\n💡 *Use /planilha para baixar o detalhamento completo em Excel.*"

    # Tenta gerar o gráfico de pizza visual
    try:
        fig, ax = plt.subplots(figsize=(6, 6))
        cores = ["#FF6B8B", "#4ECDC4", "#45B7D1", "#FFA07A", "#98D8C8", "#D4A5A5", "#F7D794", "#778BEB"]

        labels = list(categorias.keys())
        valores = list(categorias.values())

        wedges, texts, autotexts = ax.pie(
            valores,
            labels=labels,
            autopct="%1.1f%%",
            startangle=140,
            colors=cores[:len(valores)],
            wedgeprops=dict(width=0.6, edgecolor="white", linewidth=2)  # Formato Donut moderno
        )

        for autotext in autotexts:
            autotext.set_color("black")
            autotext.set_fontsize(9)
            autotext.set_weight("bold")

        ax.set_title("Distribuição de Gastos do Mês", fontsize=14, fontweight="bold", pad=20)
        plt.tight_layout()

        buf = io.BytesIO()
        plt.savefig(buf, format="png", bbox_inches="tight", dpi=160)
        plt.close(fig)
        buf.seek(0)

        await update.message.reply_photo(photo=buf, caption=mensagem, parse_mode="Markdown")
    except Exception as e:
        logging.error(f"Erro ao gerar gráfico: {e}")
        # Se der qualquer erro no gráfico, envia como mensagem de texto
        await update.message.reply_text(mensagem, parse_mode="Markdown")

# ↩️ COMANDO /desfazer (REMOVE O ÚLTIMO GASTO REGISTRADO)
async def desfazer(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = str(update.effective_user.id)
    todos_dados = carregar_todos_dados()

    if user_id not in todos_dados or not todos_dados[user_id]["gastos"]:
        await update.message.reply_text("📭 Você não tem nenhum gasto para desfazer!")
        return

    dados_usuario = todos_dados[user_id]
    ultimo_gasto = dados_usuario["gastos"].pop()

    # Devolve o valor ao saldo atual
    valor_devolvido = ultimo_gasto["valor"]
    dados_usuario["saldo_atual"] += valor_devolvido

    todos_dados[user_id] = dados_usuario
    salvar_usuario(user_id, dados_usuario)

    await update.message.reply_text(
        "↩️ **Último Gasto Desfeito com Sucesso!**\n\n"
        f"🗑️ **Removido:** {ultimo_gasto['item']} ({ultimo_gasto['categoria']})\n"
        f"💰 **Valor devolvido:** R$ {valor_devolvido:.2f}\n"
        f"📈 **Novo Saldo Atual:** R$ {dados_usuario['saldo_atual']:.2f}"
    )

# 📑 COMANDO /planilha OU /exportar (GERA ARQUIVO EXCEL .XLSX)
async def exportar_excel(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = str(update.effective_user.id)
    todos_dados = carregar_todos_dados()

    if user_id not in todos_dados or not todos_dados[user_id]["gastos"]:
        await update.message.reply_text("📭 Nenhum gasto registrado ainda para exportar!")
        return

    dados_usuario = todos_dados[user_id]
    gastos = dados_usuario["gastos"]

    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Extrato de Gastos"

        # Cores e estilos do Excel
        header_fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center")
        right_align = Alignment(horizontal="right", vertical="center")
        thin_border = Border(
            left=Side(style="thin", color="D1D5DB"),
            right=Side(style="thin", color="D1D5DB"),
            top=Side(style="thin", color="D1D5DB"),
            bottom=Side(style="thin", color="D1D5DB")
        )

        # Cabeçalho da tabela
        headers = ["Data", "Item / Descrição", "Categoria", "Valor (R$)"]
        ws.append(headers)

        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align

        # Linhas de dados
        total_gasto = 0.0
        for idx, g in enumerate(gastos, start=2):
            total_gasto += g["valor"]
            ws.append([g["data"], g["item"], g["categoria"], g["valor"]])

            ws.cell(row=idx, column=1).alignment = center_align
            ws.cell(row=idx, column=2).alignment = left_align
            ws.cell(row=idx, column=3).alignment = center_align
            val_cell = ws.cell(row=idx, column=4)
            val_cell.alignment = right_align
            val_cell.number_format = "R$ #,##0.00"

            for c in range(1, 5):
                ws.cell(row=idx, column=c).border = thin_border

        # Linha em branco e Resumo Final
        row_resumo = len(gastos) + 3

        ws.cell(row=row_resumo, column=3, value="Total Gasto:").font = Font(bold=True)
        ws.cell(row=row_resumo, column=4, value=total_gasto).font = Font(bold=True)
        ws.cell(row=row_resumo, column=4).number_format = "R$ #,##0.00"

        ws.cell(row=row_resumo + 1, column=3, value="Saldo Restante:").font = Font(bold=True, color="059669")
        ws.cell(row=row_resumo + 1, column=4, value=dados_usuario["saldo_atual"]).font = Font(bold=True, color="059669")
        ws.cell(row=row_resumo + 1, column=4).number_format = "R$ #,##0.00"

        ws.cell(row=row_resumo + 2, column=3, value="Orçamento Inicial:").font = Font(bold=True)
        ws.cell(row=row_resumo + 2, column=4, value=dados_usuario["orcamento_inicial"]).font = Font(bold=True)
        ws.cell(row=row_resumo + 2, column=4).number_format = "R$ #,##0.00"

        # Ajusta larguras das colunas
        larguras = {"A": 14, "B": 28, "C": 20, "D": 16}
        for col_letter, largura in larguras.items():
            ws.column_dimensions[col_letter].width = largura

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        nome_arquivo = f"Extrato_Porquinho_IA_{datetime.now().strftime('%m_%Y')}.xlsx"
        await update.message.reply_document(
            document=buf,
            filename=nome_arquivo,
            caption="📊 **Aqui está a sua planilha de gastos do mês pronta!**"
        )
    except Exception as e:
        logging.error(f"Erro ao gerar Excel: {e}")
        await update.message.reply_text("❌ Desculpe, tive um problema ao gerar a planilha.")

# 🔄 COMANDO /reset
async def reset(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = str(update.effective_user.id)
    todos_dados = carregar_todos_dados()

    if user_id in todos_dados:
        del todos_dados[user_id]
        deletar_usuario(user_id)

    context.user_data["configurando_renda"] = True
    await update.message.reply_text(
        "🔄 **Porquinho Reiniciado!**\n"
        "Seu histórico anterior foi zerado para dar início ao novo mês.\n\n"
        "Por favor, digite o seu **novo orçamento inicial**:"
    )

# 💡 COMANDO /ajuda (IA Consultora)
async def ajuda_ia(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = str(update.effective_user.id)
    todos_dados = carregar_todos_dados()

    pergunta_usuario = " ".join(context.args)
    if not pergunta_usuario:
        await update.message.reply_text(
            "Digite uma pergunta após o comando!\n"
            "Exemplo: `/ajuda Será que posso comprar uma blusa de 80 reais?`"
        )
        return

    if user_id not in todos_dados:
        await update.message.reply_text("Configure seu saldo inicial primeiro digitando `/start`!")
        return

    await update.message.reply_text("🤖 Analisando suas finanças... Aguarde um instante.")
    await context.bot.send_chat_action(chat_id=update.effective_chat.id, action="typing")

    dados_user = todos_dados[user_id]
    prompt_analise = (
        f"Você é o Porquinho IA, um consultor financeiro amigável e focado em ajudar o usuário.\n"
        f"Aqui estão as finanças reais do usuário atualmente:\n"
        f"- Orçamento Inicial do Mês: R$ {dados_user['orcamento_inicial']:.2f}\n"
        f"- Saldo Restante Atual: R$ {dados_user['saldo_atual']:.2f}\n"
        f"- Lista Completa de Gastos feitos até agora: {json.dumps(dados_user['gastos'])}\n\n"
        f"O usuário te fez a seguinte pergunta: '{pergunta_usuario}'.\n"
        f"Responda diretamente a pergunta dele de forma bem curta, sincera e acolhedora, "
        f"usando os dados financeiros fornecidos para justificar se ele deve ou não gastar e dê uma dica de economia."
    )

    try:
        response = ai_client.models.generate_content(
            model="gemini-2.5-flash",
            contents=prompt_analise
        )
        await update.message.reply_text(response.text, parse_mode="Markdown")
    except Exception as e:
        logging.error(f"Erro na IA consultora: {e}")
        await update.message.reply_text("❌ Desculpe, não consegui analisar suas finanças no momento.")

# 🎙️ ESCUTADOR DE ÁUDIO / MENSAGENS DE VOZ
async def escutar_audio(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = str(update.effective_user.id)
    todos_dados = carregar_todos_dados()

    if user_id not in todos_dados:
        await update.message.reply_text("Digite `/start` para começar primeiro!")
        return

    await context.bot.send_chat_action(chat_id=update.effective_chat.id, action="typing")
    mensagem_aviso = await update.message.reply_text("🎙️ *Ouvindo o seu áudio... Aguarde.*", parse_mode="Markdown")

    try:
        voice = update.message.voice or update.message.audio
        arquivo = await context.bot.get_file(voice.file_id)
        audio_bytes = await arquivo.download_as_bytearray()

        prompt_audio = (
            "Ouça o áudio do usuário e extraia o gasto mencionado: "
            "o item comprado, o valor numérico e a categoria mais adequada "
            "(ex: Alimentação, Transporte, Lazer, Pets, Contas e Utilidades, Saúde, Vestuário, Outros). "
            "Responda ESTRITAMENTE em formato JSON com as chaves 'item', 'valor' e 'categoria'."
        )

        response = ai_client.models.generate_content(
            model="gemini-2.5-flash",
            contents=[
                types.Part.from_bytes(data=bytes(audio_bytes), mime_type="audio/ogg"),
                prompt_audio
            ],
            config=types.GenerateContentConfig(response_mime_type="application/json")
        )

        dados_gasto = json.loads(response.text)
        item = dados_gasto.get("item", "Gasto por voz")
        valor_gasto = float(dados_gasto.get("valor", 0.0))
        categoria = dados_gasto.get("categoria", "Outros")

        if valor_gasto <= 0:
            await mensagem_aviso.edit_text("❌ Não consegui identificar o valor do gasto no seu áudio.")
            return

        texto_resposta = registrar_novo_gasto(todos_dados, user_id, item, valor_gasto, categoria, canal="Áudio 🎙️")
        await mensagem_aviso.edit_text(texto_resposta, parse_mode="Markdown")

    except Exception as e:
        logging.error(f"Erro ao processar áudio: {e}")
        await mensagem_aviso.edit_text("❌ Tive um problema ao processar seu áudio. Tente novamente ou digite em texto.")

# 📸 ESCUTADOR DE FOTOS (COMPROVANTES, CUPONS, NOTAS FISCAIS, PIX)
async def escutar_foto(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = str(update.effective_user.id)
    todos_dados = carregar_todos_dados()

    if user_id not in todos_dados:
        await update.message.reply_text("Digite `/start` para começar primeiro!")
        return

    await context.bot.send_chat_action(chat_id=update.effective_chat.id, action="typing")
    mensagem_aviso = await update.message.reply_text("🔍 *Analisando o comprovante/cupom fiscal...*", parse_mode="Markdown")

    try:
        # Pega a melhor resolução da foto enviada
        foto = update.message.photo[-1]
        arquivo = await context.bot.get_file(foto.file_id)
        foto_bytes = await arquivo.download_as_bytearray()

        legenda = update.message.caption or ""
        prompt_foto = (
            "Você é um assistente financeiro inteligente. Analise esta imagem, que pode ser uma nota fiscal, "
            "cupom fiscal, comprovante de transferência Pix ou recibo. "
            "Identifique o nome do estabelecimento ou item principal, o valor TOTAL pago e defina a categoria mais adequada "
            "(ex: Alimentação, Supermercado, Transporte, Farmácia, Saúde, Lazer, Contas e Utilidades, Outros). "
            f"{'Contexto adicional que o usuário digitou: ' + legenda if legenda else ''} "
            "Responda ESTRITAMENTE em formato JSON com as chaves 'item', 'valor' e 'categoria'."
        )

        response = ai_client.models.generate_content(
            model="gemini-2.5-flash",
            contents=[
                types.Part.from_bytes(data=bytes(foto_bytes), mime_type="image/jpeg"),
                prompt_foto
            ],
            config=types.GenerateContentConfig(response_mime_type="application/json")
        )

        dados_gasto = json.loads(response.text)
        item = dados_gasto.get("item", "Comprovante")
        valor_gasto = float(dados_gasto.get("valor", 0.0))
        categoria = dados_gasto.get("categoria", "Outros")

        if valor_gasto <= 0:
            await mensagem_aviso.edit_text("❌ Não consegui identificar o valor total nesta foto/comprovante.")
            return

        texto_resposta = registrar_novo_gasto(todos_dados, user_id, item, valor_gasto, categoria, canal="Foto 🧾")
        await mensagem_aviso.edit_text(texto_resposta, parse_mode="Markdown")

    except Exception as e:
        logging.error(f"Erro ao processar foto: {e}")
        await mensagem_aviso.edit_text("❌ Desculpe, não consegui ler esse comprovante. Tente uma foto mais nítida.")

# 📩 ESCUTADOR DE MENSAGENS DE TEXTO (GASTOS EM TEXTO)
async def escutar_mensagens(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = str(update.effective_user.id)
    texto_usuario = update.message.text
    todos_dados = carregar_todos_dados()

    if context.user_data.get("configurando_renda"):
        try:
            valor_renda = float(texto_usuario.replace("R$", "").replace(",", ".").strip())
            todos_dados[user_id] = {
                "orcamento_inicial": valor_renda,
                "saldo_atual": valor_renda,
                "gastos": []
            }
            salvar_usuario(user_id, todos_dados[user_id])
            context.user_data["configurando_renda"] = False
            await update.message.reply_text(
                f"🎉 Perfeito! Seu orçamento de **R$ {valor_renda:.2f}** foi configurado.\n"
                "Sempre que gastar algo, basta me enviar aqui por texto, áudio ou foto do cupom!"
            )
            return
        except ValueError:
            await update.message.reply_text("❌ Digite apenas o número, ex: `704` ou `1200`.")
            return

    if user_id not in todos_dados:
        await update.message.reply_text("Digite `/start` para começar!")
        return

    await context.bot.send_chat_action(chat_id=update.effective_chat.id, action="typing")

    prompt_ia = (
        f"Analise o seguinte gasto enviado pelo usuário: '{texto_usuario}'. "
        "Extraia o item comprado (corrija maiúsculas se necessário), o valor numérico e defina a categoria mais adequada "
        "(ex: Alimentação, Transporte, Lazer, Pets, Contas e Utilidades, Saúde, Vestuário, Outros). "
        "Responda ESTRITAMENTE em formato JSON com as chaves 'item', 'valor' e 'categoria'."
    )

    try:
        response = ai_client.models.generate_content(
            model="gemini-2.5-flash",
            contents=prompt_ia,
            config=types.GenerateContentConfig(response_mime_type="application/json")
        )

        dados_gasto = json.loads(response.text)
        item = dados_gasto.get("item", "Não identificado")
        valor_gasto = float(dados_gasto.get("valor", 0.00))
        categoria = dados_gasto.get("categoria", "Outros")

        resposta_formatada = registrar_novo_gasto(todos_dados, user_id, item, valor_gasto, categoria, canal="Texto 📝")
        await update.message.reply_text(resposta_formatada, parse_mode="Markdown")

    except Exception as e:
        logging.error(f"Erro ao processar: {e}")
        await update.message.reply_text("❌ Desculpe, tive um problema ao calcular o seu gasto.")

def main():
    if not TELEGRAM_TOKEN:
        print("❌ ERRO: A variável TELEGRAM_TOKEN não foi encontrada no arquivo .env!")
        return
    if not GEMINI_API_KEY:
        print("❌ ERRO: A variável GEMINI_API_KEY não foi encontrada no arquivo .env!")
        return

    print("🐷 Porquinho IA Premium rodando...")

    # Inicia servidor web em background para o Render e UptimeRobot
    iniciar_servidor_web()

    application = Application.builder().token(TELEGRAM_TOKEN).build()

    # Comandos
    application.add_handler(CommandHandler("start", start))
    application.add_handler(CommandHandler("extrato", extrato))
    application.add_handler(CommandHandler("desfazer", desfazer))
    application.add_handler(CommandHandler("planilha", exportar_excel))
    application.add_handler(CommandHandler("exportar", exportar_excel))
    application.add_handler(CommandHandler("reset", reset))
    application.add_handler(CommandHandler("ajuda", ajuda_ia))

    # Manipuladores de Mídia e Mensagens
    application.add_handler(MessageHandler(filters.VOICE | filters.AUDIO, escutar_audio))
    application.add_handler(MessageHandler(filters.PHOTO, escutar_foto))
    application.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, escutar_mensagens))

    application.run_polling()

if __name__ == "__main__":
    main()